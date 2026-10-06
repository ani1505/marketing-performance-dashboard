from __future__ import annotations

import os
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from filelock import FileLock
from openpyxl import load_workbook


MASTER_SHEET = "MasterData"
COMPANIES_SHEET = "Companies"
CURRENCY_SHEET = "Currency selector"

INDIVIDUAL_COLUMNS = [
    "Platform",
    "S.No",
    "Campaign Name",
    "Campaign Status",
    "Region",
    "Start Date",
    "End Date",
    "Currency",
    "Daily Budget",
    "Impressions",
    "Clicks",
    "Avg. CPC",
    "Total Leads",
    "Cost Per Lead",
    "Total Spend",
    "Notes",
]

MASTER_COLUMNS = [
    "Company ID",
    "Company",
    "Platform",
    "S.No",
    "Campaign Name",
    "Campaign Status",
    "Region",
    "Report Date",
    "Start Date",
    "End Date",
    "Currency",
    "Daily Budget",
    "Impressions",
    "Clicks",
    "Avg. CPC",
    "Total Leads",
    "Cost Per Lead",
    "Total Spend",
    "Notes",
]

# Explicit business rules supplied for the dashboard.
COMPANY_CURRENCY_RULES: dict[str, list[str]] = {
    "Wolfable": ["INR", "GBP", "USD"],
    "Wellbeing": ["USD"],
    "Keecare": ["GBP"],
    "Cricket Centurion": ["USD"],
    "M360": ["USD"],
    "Elite Valuation": ["INR"],
    "Musk": ["INR"],
}


@dataclass
class SaveResult:
    company: str
    company_id: str
    s_no: int
    action: str
    backup_path: Path


class WorkbookStore:
    """Read/write helper that keeps MasterData and the individual company sheet in sync."""

    def __init__(self, workbook_path: str | Path):
        self.path = Path(workbook_path).expanduser().resolve()
        self.lock_path = Path(str(self.path) + ".lock")
        self.backup_dir = self.path.parent / "backups"

    def ensure_exists(self) -> None:
        if not self.path.exists():
            raise FileNotFoundError(
                f"Workbook not found: {self.path}. Set DASHBOARD_WORKBOOK or place "
                "wolfable_master.xlsx beside app.py."
            )

    @staticmethod
    def _clean_df(df: pd.DataFrame) -> pd.DataFrame:
        df = df.copy()
        df = df.dropna(how="all")
        df.columns = [str(c).strip() for c in df.columns]
        return df

    def read_sheet(self, sheet_name: str) -> pd.DataFrame:
        self.ensure_exists()
        return self._clean_df(pd.read_excel(self.path, sheet_name=sheet_name, engine="openpyxl"))

    def load_dashboard_data(self) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
        master = self.read_sheet(MASTER_SHEET)
        companies = self.read_sheet(COMPANIES_SHEET)
        currencies = self.read_sheet(CURRENCY_SHEET)

        for col in ["Report Date", "Start Date", "End Date"]:
            if col in master.columns:
                master[col] = pd.to_datetime(master[col], errors="coerce")

        numeric_cols = [
            "S.No",
            "Daily Budget",
            "Impressions",
            "Clicks",
            "Avg. CPC",
            "Total Leads",
            "Cost Per Lead",
            "Total Spend",
        ]
        for col in numeric_cols:
            if col in master.columns:
                master[col] = pd.to_numeric(master[col], errors="coerce")

        return master, companies, currencies

    def get_company_rows(self, company: str) -> pd.DataFrame:
        master = self.read_sheet(MASTER_SHEET)
        if "Company" not in master.columns:
            return pd.DataFrame(columns=MASTER_COLUMNS)
        out = master[master["Company"].astype(str).str.strip() == company].copy()
        if "S.No" in out.columns:
            out["S.No"] = pd.to_numeric(out["S.No"], errors="coerce")
            out = out.sort_values("S.No")
        return out

    def company_lookup(self) -> dict[str, str]:
        companies = self.read_sheet(COMPANIES_SHEET)
        return {
            str(row["Company Name"]).strip(): str(row["Company ID"]).strip()
            for _, row in companies.iterrows()
            if pd.notna(row.get("Company Name")) and pd.notna(row.get("Company ID"))
        }

    def currency_symbols(self) -> dict[str, str]:
        df = self.read_sheet(CURRENCY_SHEET)
        return {
            str(row["Currency"]).strip(): str(row["Symbol"]).strip()
            for _, row in df.iterrows()
            if pd.notna(row.get("Currency")) and pd.notna(row.get("Symbol"))
        }

    def allowed_currencies(self, company: str) -> list[str]:
        if company in COMPANY_CURRENCY_RULES:
            return COMPANY_CURRENCY_RULES[company]
        companies = self.read_sheet(COMPANIES_SHEET)
        match = companies[companies["Company Name"].astype(str).str.strip() == company]
        if match.empty:
            return list(self.currency_symbols().keys())
        native = str(match.iloc[0].get("Native Currency", "")).strip()
        return [native] if native and native.lower() != "mixed" else list(self.currency_symbols().keys())

    @staticmethod
    def _headers(ws) -> dict[str, int]:
        return {
            str(cell.value).strip(): cell.column
            for cell in ws[1]
            if cell.value is not None and str(cell.value).strip()
        }

    @staticmethod
    def _find_row(ws, header_map: dict[str, int], conditions: dict[str, Any]) -> int | None:
        for row in range(2, ws.max_row + 1):
            matched = True
            for header, expected in conditions.items():
                col = header_map.get(header)
                if col is None:
                    matched = False
                    break
                actual = ws.cell(row=row, column=col).value
                if str(actual).strip() != str(expected).strip():
                    matched = False
                    break
            if matched:
                return row
        return None

    @staticmethod
    def _copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
        if source_row < 2 or source_row > ws.max_row:
            return
        for col in range(1, max_col + 1):
            src = ws.cell(source_row, col)
            dst = ws.cell(target_row, col)
            if src.has_style:
                dst._style = copy(src._style)
            if src.number_format:
                dst.number_format = src.number_format
            if src.font:
                dst.font = copy(src.font)
            if src.fill:
                dst.fill = copy(src.fill)
            if src.border:
                dst.border = copy(src.border)
            if src.alignment:
                dst.alignment = copy(src.alignment)
            if src.protection:
                dst.protection = copy(src.protection)
        if source_row in ws.row_dimensions:
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

    @staticmethod
    def _write_record(ws, row: int, header_map: dict[str, int], record: dict[str, Any]) -> None:
        for key, value in record.items():
            col = header_map.get(key)
            if col is not None:
                ws.cell(row=row, column=col).value = value

    @staticmethod
    def _excel_date(value: Any) -> Any:
        if value is None or value == "":
            return None
        if isinstance(value, pd.Timestamp):
            return value.to_pydatetime()
        if isinstance(value, datetime):
            return value
        if isinstance(value, date):
            return datetime.combine(value, datetime.min.time())
        parsed = pd.to_datetime(value, errors="coerce")
        return None if pd.isna(parsed) else parsed.to_pydatetime()

    def _create_backup(self) -> Path:
        self.backup_dir.mkdir(exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        backup = self.backup_dir / f"{self.path.stem}_{stamp}{self.path.suffix}"
        shutil.copy2(self.path, backup)
        backups = sorted(self.backup_dir.glob(f"{self.path.stem}_*{self.path.suffix}"))
        for old in backups[:-20]:
            try:
                old.unlink()
            except OSError:
                pass
        return backup

    def _atomic_save(self, wb) -> None:
        temp = self.path.with_name(f".{self.path.stem}.tmp{self.path.suffix}")
        try:
            wb.save(temp)
            os.replace(temp, self.path)
        finally:
            if temp.exists():
                try:
                    temp.unlink()
                except OSError:
                    pass

    def _next_sno(self, wb, company: str, company_id: str) -> int:
        values: list[int] = []

        if company in wb.sheetnames:
            ws = wb[company]
            hm = self._headers(ws)
            col = hm.get("S.No")
            if col:
                for row in range(2, ws.max_row + 1):
                    val = ws.cell(row, col).value
                    try:
                        values.append(int(float(val)))
                    except (TypeError, ValueError):
                        pass

        ws_master = wb[MASTER_SHEET]
        hm = self._headers(ws_master)
        sno_col = hm.get("S.No")
        cid_col = hm.get("Company ID")
        if sno_col and cid_col:
            for row in range(2, ws_master.max_row + 1):
                if str(ws_master.cell(row, cid_col).value).strip() == company_id:
                    try:
                        values.append(int(float(ws_master.cell(row, sno_col).value)))
                    except (TypeError, ValueError):
                        pass
        return max(values, default=0) + 1

    def save_record(
        self,
        company: str,
        record: dict[str, Any],
        report_date: date | datetime,
        s_no: int | None = None,
    ) -> SaveResult:
        self.ensure_exists()
        with FileLock(str(self.lock_path), timeout=15):
            backup = self._create_backup()
            wb = load_workbook(self.path)
            try:
                if company not in wb.sheetnames:
                    raise ValueError(f"Individual company sheet '{company}' does not exist.")
                if MASTER_SHEET not in wb.sheetnames or COMPANIES_SHEET not in wb.sheetnames:
                    raise ValueError("Workbook must contain MasterData and Companies sheets.")

                # Company ID is read from the actual workbook to avoid stale app state.
                ws_companies = wb[COMPANIES_SHEET]
                ch = self._headers(ws_companies)
                id_col = ch.get("Company ID")
                name_col = ch.get("Company Name")
                company_id = None
                if id_col and name_col:
                    for row in range(2, ws_companies.max_row + 1):
                        if str(ws_companies.cell(row, name_col).value).strip() == company:
                            company_id = str(ws_companies.cell(row, id_col).value).strip()
                            break
                if not company_id:
                    raise ValueError(f"No Company ID found for {company!r} in Companies sheet.")

                allowed = self.allowed_currencies(company)
                currency = str(record.get("Currency", "")).strip()
                if currency not in allowed:
                    raise ValueError(f"{currency} is not an allowed currency for {company}. Allowed: {allowed}")

                if s_no is None:
                    s_no = self._next_sno(wb, company, company_id)
                    action = "created"
                else:
                    s_no = int(s_no)
                    action = "updated"

                clicks = float(record.get("Clicks") or 0)
                leads = float(record.get("Total Leads") or 0)
                spend = float(record.get("Total Spend") or 0)
                avg_cpc = spend / clicks if clicks else 0.0
                cpl = spend / leads if leads else 0.0

                individual_record = {
                    "Platform": record.get("Platform"),
                    "S.No": s_no,
                    "Campaign Name": record.get("Campaign Name"),
                    "Campaign Status": record.get("Campaign Status"),
                    "Region": record.get("Region"),
                    "Start Date": self._excel_date(record.get("Start Date")),
                    "End Date": self._excel_date(record.get("End Date")),
                    "Currency": currency,
                    "Daily Budget": float(record.get("Daily Budget") or 0),
                    "Impressions": int(record.get("Impressions") or 0),
                    "Clicks": int(record.get("Clicks") or 0),
                    "Avg. CPC": round(avg_cpc, 2),
                    "Total Leads": int(record.get("Total Leads") or 0),
                    "Cost Per Lead": round(cpl, 2),
                    "Total Spend": round(spend, 2),
                    "Notes": record.get("Notes") or None,
                }

                ws_company = wb[company]
                ih = self._headers(ws_company)
                missing = [c for c in INDIVIDUAL_COLUMNS if c not in ih]
                if missing:
                    raise ValueError(f"Sheet {company} is missing columns: {missing}")
                row_ind = self._find_row(ws_company, ih, {"S.No": s_no})
                if row_ind is None:
                    row_ind = ws_company.max_row + 1
                    self._copy_row_style(ws_company, max(2, row_ind - 1), row_ind, len(INDIVIDUAL_COLUMNS))
                self._write_record(ws_company, row_ind, ih, individual_record)

                master_record = {
                    "Company ID": company_id,
                    "Company": company,
                    "Platform": individual_record["Platform"],
                    "S.No": s_no,
                    "Campaign Name": individual_record["Campaign Name"],
                    "Campaign Status": individual_record["Campaign Status"],
                    "Region": individual_record["Region"],
                    "Report Date": self._excel_date(report_date),
                    "Start Date": individual_record["Start Date"],
                    "End Date": individual_record["End Date"],
                    "Currency": currency,
                    "Daily Budget": individual_record["Daily Budget"],
                    "Impressions": individual_record["Impressions"],
                    "Clicks": individual_record["Clicks"],
                    "Avg. CPC": individual_record["Avg. CPC"],
                    "Total Leads": individual_record["Total Leads"],
                    "Cost Per Lead": individual_record["Cost Per Lead"],
                    "Total Spend": individual_record["Total Spend"],
                    "Notes": individual_record["Notes"],
                }

                ws_master = wb[MASTER_SHEET]
                mh = self._headers(ws_master)
                missing_master = [c for c in MASTER_COLUMNS if c not in mh]
                if missing_master:
                    raise ValueError(f"MasterData is missing columns: {missing_master}")
                row_master = self._find_row(
                    ws_master,
                    mh,
                    {"Company ID": company_id, "S.No": s_no},
                )
                if row_master is None:
                    row_master = ws_master.max_row + 1
                    self._copy_row_style(ws_master, max(2, row_master - 1), row_master, len(MASTER_COLUMNS))
                self._write_record(ws_master, row_master, mh, master_record)

                # Excel-friendly date formats on both sheets.
                for header in ["Start Date", "End Date"]:
                    ws_company.cell(row_ind, ih[header]).number_format = "dd-mmm-yyyy"
                for header in ["Report Date", "Start Date", "End Date"]:
                    ws_master.cell(row_master, mh[header]).number_format = "dd-mmm-yyyy"

                self._atomic_save(wb)
            finally:
                wb.close()

            return SaveResult(company, company_id, int(s_no), action, backup)

    def delete_record(self, company: str, s_no: int) -> SaveResult:
        self.ensure_exists()
        with FileLock(str(self.lock_path), timeout=15):
            backup = self._create_backup()
            wb = load_workbook(self.path)
            try:
                companies = self.company_lookup()
                company_id = companies.get(company)
                if not company_id:
                    raise ValueError(f"Company ID not found for {company}")

                deleted = False
                if company in wb.sheetnames:
                    ws = wb[company]
                    hm = self._headers(ws)
                    row = self._find_row(ws, hm, {"S.No": int(s_no)})
                    if row:
                        ws.delete_rows(row, 1)
                        deleted = True

                ws_master = wb[MASTER_SHEET]
                hm = self._headers(ws_master)
                rows_to_delete = []
                for row in range(2, ws_master.max_row + 1):
                    cid = str(ws_master.cell(row, hm["Company ID"]).value).strip()
                    sno = str(ws_master.cell(row, hm["S.No"]).value).strip()
                    if cid == company_id and sno in {str(int(s_no)), str(float(s_no))}:
                        rows_to_delete.append(row)
                for row in reversed(rows_to_delete):
                    ws_master.delete_rows(row, 1)
                    deleted = True

                if not deleted:
                    raise ValueError(f"Record {company} / S.No {s_no} was not found.")

                self._atomic_save(wb)
            finally:
                wb.close()

            return SaveResult(company, company_id, int(s_no), "deleted", backup)
